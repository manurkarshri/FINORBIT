from pathlib import Path
import sys

from docx import Document


def escape(text: str) -> str:
    return text.replace("\\", "\\\\").replace("|", "\\|").strip()


def main() -> None:
    source = Path(sys.argv[1])
    target = Path(sys.argv[2])
    document = Document(source)
    lines: list[str] = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue

        style = paragraph.style.name if paragraph.style else ""
        if style.startswith("Heading "):
            level = int(style.split()[-1])
            lines.extend([f"{'#' * level} {text}", ""])
        elif style.startswith("List Bullet"):
            lines.extend([f"- {text}", ""])
        elif style.startswith("List Number"):
            lines.extend([f"1. {text}", ""])
        else:
            lines.extend([text, ""])

    for table in document.tables:
        rows = [[escape(cell.text.replace("\n", "<br>")) for cell in row.cells] for row in table.rows]
        if not rows:
            continue
        lines.append("| " + " | ".join(rows[0]) + " |")
        lines.append("| " + " | ".join("---" for _ in rows[0]) + " |")
        lines.extend("| " + " | ".join(row) + " |" for row in rows[1:])
        lines.append("")

    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text("\n".join(lines).rstrip() + "\n", encoding="utf-8")


if __name__ == "__main__":
    main()
